"""自訂名單 (stock_SB) 乖離率監控 Excel"""
import os
import time
import random

import pandas as pd
import matplotlib.pyplot as plt

from config import *
from stock_ai.data_fetch import read_price_csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_custom_list_bias_excel(stock_sb_dir, source_dir, output_dir, today_str, ticker_map, industry_map):
    """讀取 stock_SB 資料夾下的所有 txt，取得資料、計算乖離率並輸出包含出現紀錄與產業分類的 Excel"""
    print(f"\n📊 開始產生【自訂名單乖離率監控表】 (讀取 {stock_sb_dir} 目錄)...")
    if not os.path.exists(stock_sb_dir):
        print(f"⚠️ 找不到 {stock_sb_dir} 資料夾，略過此步驟。(若需要產生，請在同目錄下建立此資料夾並放入 txt)")
        return

    stock_appearances = {}
    for filename in os.listdir(stock_sb_dir):
        if filename.endswith(".txt"):
            date_name = filename.replace(".txt", "")
            file_path = os.path.join(stock_sb_dir, filename)
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                for line in f:
                    s_id = line.strip()
                    if s_id:
                        if s_id not in stock_appearances:
                            stock_appearances[s_id] = set()
                        stock_appearances[s_id].add(date_name)

    if not stock_appearances:
        print(f"⚠️ {stock_sb_dir} 資料夾內沒有有效的股票代碼，略過此步驟。")
        return

    id_to_ticker = {}
    for ticker in ticker_map.keys():
        s_id = ticker.split('.')[0]
        id_to_ticker[s_id] = ticker
        
    data_list = []
    for s_id, dates_set in stock_appearances.items():
        ticker = id_to_ticker.get(s_id)
        if not ticker:
            if os.path.exists(os.path.join(source_dir, f"{s_id}.TW.csv")):
                ticker = f"{s_id}.TW"
            elif os.path.exists(os.path.join(source_dir, f"{s_id}.TWO.csv")):
                ticker = f"{s_id}.TWO"
            else:
                continue
                
        file_path = os.path.join(source_dir, f"{ticker}.csv")
        if not os.path.exists(file_path):
            continue
            
        try:
            df = read_price_csv(file_path)
            df.index = pd.to_datetime(df.index, errors='coerce')
            df = df[df.index.notnull()].sort_index()
            
            if 'Close' not in df.columns: continue
            df['Close'] = pd.to_numeric(df['Close'], errors='coerce')
            df = df.dropna(subset=['Close'])
            if df.empty or len(df) < 5: continue
            
            df['5MA'] = df['Close'].rolling(window=5).mean()
            df['10MA'] = df['Close'].rolling(window=10).mean()
            df['20MA'] = df['Close'].rolling(window=20).mean()
            df['60MA'] = df['Close'].rolling(window=60).mean()

            df['5日乖離率'] = ((df['Close'] - df['5MA']) / df['5MA']) * 100
            df['10日乖離率'] = ((df['Close'] - df['10MA']) / df['10MA']) * 100
            df['20日乖離率'] = ((df['Close'] - df['20MA']) / df['20MA']) * 100
            df['60日乖離率'] = ((df['Close'] - df['60MA']) / df['60MA']) * 100
            
            latest_data = df.iloc[-1]
            date_str = df.index[-1].strftime('%Y-%m-%d')
            
            stock_name = ticker_map.get(ticker, "未知")
            industry_name = industry_map.get(ticker, "未分類")
            
            sorted_dates = sorted(list(dates_set), reverse=True)
            formatted_dates = []
            for d in sorted_dates:
                if len(d) == 8 and d.isdigit():
                    month = int(d[4:6])
                    day = int(d[6:8])
                    formatted_dates.append(f"{month}/{day}")
                else:
                    formatted_dates.append(d)
            appear_records = ", ".join(formatted_dates)

            data_list.append({
                '代碼': s_id, '名稱': stock_name, '產業分類': industry_name, '出現紀錄': appear_records, '日期': date_str,
                '目前股價': float(latest_data['Close']),
                '5日均線': float(latest_data['5MA']) if pd.notna(latest_data['5MA']) else None,
                '10日均線': float(latest_data['10MA']) if pd.notna(latest_data['10MA']) else None,
                '20日均線': float(latest_data['20MA']) if pd.notna(latest_data['20MA']) else None,
                '60日均線': float(latest_data['60MA']) if pd.notna(latest_data['60MA']) else None,
                '5日乖離率': float(latest_data['5日乖離率']) if pd.notna(latest_data['5日乖離率']) else None,
                '10日乖離率': float(latest_data['10日乖離率']) if pd.notna(latest_data['10日乖離率']) else None,
                '20日乖離率': float(latest_data['20日乖離率']) if pd.notna(latest_data['20日乖離率']) else None,
                '60日乖離率': float(latest_data['60日乖離率']) if pd.notna(latest_data['60日乖離率']) else None,
            })
        except Exception:
            continue
            
    if not data_list:
        print("⚠️ 沒有足夠的有效數據可供產出乖離率報表。")
        return
        
    filename = os.path.join(output_dir, f"自訂名單_技術指標監控表_{today_str}.xlsx")
    wb = Workbook()
    
    header_fill = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")
    zebra_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    positive_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    negative_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
    
    font_header = Font(name="Microsoft JhengHei", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Microsoft JhengHei", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3')
    )

    def format_worksheet(sheet, header_list, percent_cols):
        for col_num, header in enumerate(header_list, 1):
            cell = sheet.cell(row=1, column=col_num)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        sheet.row_dimensions[1].height = 28
        
        for row_idx in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(header_list) + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                cell.font = font_body
                cell.border = thin_border
                if row_idx % 2 == 1: cell.fill = zebra_fill
                
                if col_idx <= 4: cell.alignment = Alignment(horizontal="center", vertical="center")
                else: cell.alignment = Alignment(horizontal="right", vertical="center")
                
                if col_idx in percent_cols:
                    cell.number_format = '+0.00%;-0.00%;0.00%'
                    val = cell.value
                    if val is not None and isinstance(val, (int, float)):
                        cell.value = val / 100
                        if val > 5: cell.fill = positive_fill
                        elif val < -5: cell.fill = negative_fill
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'

        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    byte_len = len(val_str.encode('utf-8'))
                    if byte_len > max_len: max_len = byte_len
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 1
    ws1 = wb.active
    ws1.title = "技術指標監控"
    ws1.views.sheetView[0].showGridLines = True
    headers1 = ['代碼', '名稱', '產業分類', '出現紀錄', '資料日期', '目前股價', '5日均線', '10日均線', '20日均線', '60日均線', '乖離率(5日)', '乖離率(10日)', '乖離率(20日)', '乖離率(60日)']
    ws1.append(headers1)
    for data in data_list:
        ws1.append([
            data['代碼'], data['名稱'], data['產業分類'], data['出現紀錄'], data['日期'], data['目前股價'],
            data['5日均線'], data['10日均線'], data['20日均線'], data['60日均線'],
            data['5日乖離率'], data['10日乖離率'], data['20日乖離率'], data['60日乖離率']
        ])
    format_worksheet(ws1, headers1, percent_cols=[11, 12, 13, 14])

    # Sheet 2
    ws2 = wb.create_sheet(title="進出場機會")
    ws2.views.sheetView[0].showGridLines = True
    headers2 = ['策略分類', '代碼', '名稱', '產業分類', '目前股價', '5日乖離率', '10日乖離率', '20日乖離率']
    ws2.append(headers2)
    opp_list = []
    for d in data_list:
        if any(d[k] is None for k in ['5日均線', '10日均線', '20日均線', '60日均線', '5日乖離率', '10日乖離率', '20日乖離率']):
            continue
        b5, b10, b20 = d['5日乖離率'], d['10日乖離率'], d['20日乖離率']
        m5, m10, m20, m60 = d['5日均線'], d['10日均線'], d['20日均線'], d['60日均線']
        strat_type = ""
        if b5 > 15.0 or b10 > 20.0: strat_type = "⚠️ 逢高減碼 (乖離過大)"
        elif (m5 > m10 > m20 > m60) and (2.0 < b5 < 6.0): strat_type = "🚀 順勢動能 (多頭行進)"
        elif (m20 > m60) and (-5.0 <= b5 <= 2.0) and (b20 > 0): strat_type = "🎯 逢回佈局 (回檔有守)"
        if strat_type: opp_list.append([strat_type, d['代碼'], d['名稱'], d['產業分類'], d['目前股價'], b5, b10, b20])
    opp_list.sort(key=lambda x: x[0])
    for row in opp_list: ws2.append(row)
    format_worksheet(ws2, headers2, percent_cols=[6, 7, 8])

    # Sheet 3
    ws3 = wb.create_sheet(title="族群強弱追蹤")
    ws3.views.sheetView[0].showGridLines = True
    headers3 = ['產業分類', '追蹤檔數', '平均5日乖離率', '平均20日乖離率', '站上月線比例(強勢度)']
    ws3.append(headers3)
    ind_stats = {}
    for d in data_list:
        ind = d['產業分類']
        if not ind or ind == "nan": ind = "未分類"
        if ind not in ind_stats: ind_stats[ind] = {'count': 0, 'sum_b5': 0, 'sum_b20': 0, 'above_m20': 0}
        ind_stats[ind]['count'] += 1
        if d['5日乖離率'] is not None: ind_stats[ind]['sum_b5'] += d['5日乖離率']
        if d['20日乖離率'] is not None: ind_stats[ind]['sum_b20'] += d['20日乖離率']
        if d['目前股價'] is not None and d['20日均線'] is not None and d['目前股價'] > d['20日均線']: ind_stats[ind]['above_m20'] += 1
    ind_rows = []
    for ind, stats in ind_stats.items():
        c = stats['count']
        ind_rows.append([ind, c, stats['sum_b5']/c, stats['sum_b20']/c, (stats['above_m20']/c)*100])
    ind_rows.sort(key=lambda x: x[3], reverse=True)
    for row in ind_rows: ws3.append(row)
    format_worksheet(ws3, headers3, percent_cols=[3, 4, 5])

    wb.save(filename)
    print(f"✅ 成功匯出多重分頁監控報表至：{filename}")
